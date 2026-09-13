import logging
import os
from typing import List
from app.models.share_profile import ShareProfile

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger(__name__)

class XLSXExporter:
    @staticmethod
    def export(profiles: List[ShareProfile], output_path: str):
        if not OPENPYXL_AVAILABLE:
            logger.error("openpyxl is not installed. Cannot export to XLSX.")
            return False
            
        try:
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            wb = Workbook()
            ws = wb.active
            ws.title = "Facebook Shares"
            
            headers = [
                'Name', 'Profile URL', 'Share URL', 'Confidence', 
                'Source', 'First Seen', 'Last Seen', 'Session ID'
            ]
            
            # Write headers
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
                
            # Freeze panes
            ws.freeze_panes = "A2"
            
            # Write data
            for row, profile in enumerate(profiles, 2):
                ws.cell(row=row, column=1, value=profile.name)
                ws.cell(row=row, column=2, value=profile.profile_url or "")
                ws.cell(row=row, column=3, value=profile.share_url or "")
                ws.cell(row=row, column=4, value=profile.confidence)
                ws.cell(row=row, column=5, value=profile.source)
                ws.cell(row=row, column=6, value=profile.first_seen_at)
                ws.cell(row=row, column=7, value=profile.last_seen_at)
                ws.cell(row=row, column=8, value=profile.session_id)
                
            # Autofilter
            ws.auto_filter.ref = ws.dimensions
            
            # Adjust column widths roughly
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 50
            ws.column_dimensions['C'].width = 50
            ws.column_dimensions['D'].width = 12
            
            wb.save(output_path)
            logger.info(f"Successfully exported {len(profiles)} profiles to {output_path}")
            return True
        except Exception as e:
            logger.error(f"Failed to export XLSX to {output_path}: {e}")
            return False
